import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def add_autofill_assets():
    # ----------------------------------------------------
    # 1. LOAD SPREADSHEET AND ADD PRODUCT-SUPPLIER MAP
    # ----------------------------------------------------
    print("Loading Inventory_Management_System.xlsx...")
    wb = openpyxl.load_workbook("Inventory_Management_System.xlsx")
    
    # Get sheets
    ws_refs = wb["ReferenceLists"]
    ws_prod = wb["Product Master"]
    
    # Define mapping data headers
    ws_refs["E1"] = "Product ID"
    ws_refs["F1"] = "Supplier Name"
    ws_refs["E1"].font = Font(name="Segoe UI", bold=True)
    ws_refs["F1"].font = Font(name="Segoe UI", bold=True)
    
    # Multi-supplier mappings (Product ID -> List of Suppliers)
    multi_suppliers = {
        "PRD-00001": ["Apex Distributors", "Prime Logistics"],
        "PRD-00002": ["Apex Distributors", "Summit Wholesale"],
        "PRD-00003": ["Global Trade Co.", "Apex Distributors"],
        "PRD-00006": ["Summit Wholesale", "Prime Logistics"],
        "PRD-00011": ["Summit Wholesale", "Global Trade Co."],
        "PRD-00015": ["Summit Wholesale", "Prime Logistics"],
        "PRD-00021": ["Prime Logistics", "Apex Distributors"],
        "PRD-00031": ["BioPharma Supplies", "Summit Wholesale"],
        "PRD-00041": ["Global Trade Co.", "Prime Logistics"]
    }
    
    # Read all products from Product Master sheet to populate the map
    # Row 7 is where data starts.
    r = 7
    map_row = 2
    while True:
        prod_id_formula = ws_prod.cell(row=r, column=1).value
        # If it's empty, we stop
        if prod_id_formula is None:
            break
        
        # We can construct the product ID dynamically based on row
        prod_id = f"PRD-{r-6:05d}"
        primary_sup = ws_prod.cell(row=r, column=5).value
        
        # If it is in multi-supplier dictionary, write both rows
        if prod_id in multi_suppliers:
            for sup in multi_suppliers[prod_id]:
                ws_refs.cell(row=map_row, column=5, value=prod_id)
                ws_refs.cell(row=map_row, column=6, value=sup)
                map_row += 1
        else:
            # Write primary supplier only
            ws_refs.cell(row=map_row, column=5, value=prod_id)
            ws_refs.cell(row=map_row, column=6, value=primary_sup)
            map_row += 1
            
        r += 1

    # Save Excel file
    wb.save("Inventory_Management_System.xlsx")
    print(f"Relational Product-Supplier Map written successfully (Total {map_row-1} rows). Saved spreadsheet.")

    # ----------------------------------------------------
    # 2. GENERATE EXCEL VBA FILE CONTENTS
    # ----------------------------------------------------
    
    vba_purch = """' =========================================================================
' SHEET CODE FOR: Purchase Entry
' INSTRUCTIONS: Double click "Sheet3 (Purchase Entry)" in VBA Project explorer 
'               and paste this entire block of code.
' =========================================================================
Private Sub Worksheet_Change(ByVal Target As Range)
    Dim KeyCells As Range
    ' Monitor column C (Product ID), rows 7 onwards
    Set KeyCells = Range("C7:C10000")
    
    Dim TargetInterest As Range
    Set TargetInterest = Intersect(Target, KeyCells)
    
    If Not TargetInterest Is Nothing Then
        Dim Cell As Range
        Dim r As Long
        Dim prodId As String
        Dim ws_prod As Worksheet
        Dim ws_refs As Worksheet
        Dim matchIdx As Variant
        
        On Error GoTo ErrHandler
        Application.EnableEvents = False
        
        Set ws_prod = Worksheets("Product Master")
        Set ws_refs = Worksheets("ReferenceLists")
        
        For Each Cell In TargetInterest
            r = Cell.Row
            prodId = Trim(Cell.Value)
            
            If prodId = "" Then
                ' Clear cells if Product ID is deleted
                Cells(r, 4).Value = "" ' Product Name
                Cells(r, 5).Value = "" ' Supplier Name
                Cells(r, 7).Value = "" ' Cost Price
                Cells(r, 5).Validation.Delete
            Else
                ' Search for Product ID in Product Master Column A (Columns(1))
                matchIdx = Application.Match(prodId, ws_prod.Columns(1), 0)
                
                If IsError(matchIdx) Then
                    MsgBox "Product ID '" & prodId & "' was not found in the Product Master registry!" & vbCrLf & _
                           "Please register this product in the Product Master sheet first.", vbExclamation + vbOKOnly, "Invalid Product ID"
                    
                    ' Clear the row details
                    Cell.Value = ""
                    Cells(r, 4).Value = ""
                    Cells(r, 5).Value = ""
                    Cells(r, 7).Value = ""
                    Cells(r, 5).Validation.Delete
                Else
                    ' Populate Product Name (Column B / 2) & Cost Price (Column H / 8)
                    Cells(r, 4).Value = ws_prod.Cells(matchIdx, 2).Value
                    Cells(r, 7).Value = ws_prod.Cells(matchIdx, 8).Value
                    
                    ' Search for all matching suppliers in ReferenceLists Column E & F
                    Dim lastRowRefs As Long
                    Dim i As Long
                    Dim suppliersList As String
                    
                    suppliersList = ""
                    lastRowRefs = ws_refs.Cells(ws_refs.Rows.Count, "E").End(xlUp).Row
                    
                    For i = 2 To lastRowRefs
                        If ws_refs.Cells(i, 5).Value = prodId Then
                            If suppliersList = "" Then
                                suppliersList = ws_refs.Cells(i, 6).Value
                            Else
                                suppliersList = suppliersList & "," & ws_refs.Cells(i, 6).Value
                            End If
                        End If
                    Next i
                    
                    ' Fallback if no mapping exists (use primary supplier in Product Master Column E / 5)
                    If suppliersList = "" Then
                        suppliersList = ws_prod.Cells(matchIdx, 5).Value
                    End If
                    
                    ' Check if multiple suppliers exist
                    If InStr(suppliersList, ",") > 0 Then
                        ' Create Dropdown Validation for Column E (Supplier Name)
                        With Cells(r, 5).Validation
                            .Delete
                            .Add Type:=xlValidateList, AlertStyle:=xlValidAlertStop, Operator:= _
                            xlBetween, Formula1:=suppliersList
                            .IgnoreBlank = True
                            .InCellDropdown = True
                            .ShowInput = True
                            .ShowError = True
                        End With
                        
                        ' Auto-select the first supplier by default
                        Cells(r, 5).Value = Split(suppliersList, ",")(0)
                    Else
                        ' Single supplier: Clear any previous validation dropdown, write value directly
                        Cells(r, 5).Validation.Delete
                        Cells(r, 5).Value = suppliersList
                    End If
                End If
            End If
        Next Cell
        
CleanExit:
        Application.EnableEvents = True
        Exit Sub
ErrHandler:
        MsgBox "An error occurred during auto-fill: " & Err.Description, vbCritical, "Auto-Fill Error"
        Resume CleanExit
    End If
End Sub
"""

    vba_sales = """' =========================================================================
' SHEET CODE FOR: Sales Entry
' INSTRUCTIONS: Double click "Sheet4 (Sales Entry)" in VBA Project explorer 
'               and paste this entire block of code.
' =========================================================================
Private Sub Worksheet_Change(ByVal Target As Range)
    Dim KeyCells As Range
    ' Monitor column C (Product ID), rows 7 onwards
    Set KeyCells = Range("C7:C10000")
    
    Dim TargetInterest As Range
    Set TargetInterest = Intersect(Target, KeyCells)
    
    If Not TargetInterest Is Nothing Then
        Dim Cell As Range
        Dim r As Long
        Dim prodId As String
        Dim ws_prod As Worksheet
        Dim matchIdx As Variant
        
        On Error GoTo ErrHandler
        Application.EnableEvents = False
        
        Set ws_prod = Worksheets("Product Master")
        
        For Each Cell In TargetInterest
            r = Cell.Row
            prodId = Trim(Cell.Value)
            
            If prodId = "" Then
                ' Clear cells if Product ID is deleted
                Cells(r, 4).Value = "" ' Product Name
                Cells(r, 6).Value = "" ' Selling Price
            Else
                ' Search for Product ID in Product Master Column A (Columns(1))
                matchIdx = Application.Match(prodId, ws_prod.Columns(1), 0)
                
                If IsError(matchIdx) Then
                    MsgBox "Product ID '" & prodId & "' was not found in the Product Master registry!" & vbCrLf & _
                           "Please register this product in the Product Master sheet first.", vbExclamation + vbOKOnly, "Invalid Product ID"
                    
                    ' Clear the row details
                    Cell.Value = ""
                    Cells(r, 4).Value = ""
                    Cells(r, 6).Value = ""
                Else
                    ' Populate Product Name (Column B / 2) & Selling Price (Column I / 9)
                    Cells(r, 4).Value = ws_prod.Cells(matchIdx, 2).Value
                    Cells(r, 6).Value = ws_prod.Cells(matchIdx, 9).Value
                End If
            End If
        Next Cell
        
CleanExit:
        Application.EnableEvents = True
        Exit Sub
ErrHandler:
        MsgBox "An error occurred during auto-fill: " & Err.Description, vbCritical, "Auto-Fill Error"
        Resume CleanExit
    End If
End Sub
"""

    vba_inv = """' =========================================================================
' SHEET CODE FOR: Invoice
' INSTRUCTIONS: Double click "Sheet7 (Invoice)" in VBA Project explorer 
'               and paste this entire block of code.
' =========================================================================
Private Sub Worksheet_Change(ByVal Target As Range)
    Dim KeyCells As Range
    ' Monitor column B (Product ID), rows 11 to 20
    Set KeyCells = Range("B11:B20")
    
    Dim TargetInterest As Range
    Set TargetInterest = Intersect(Target, KeyCells)
    
    If Not TargetInterest Is Nothing Then
        Dim Cell As Range
        Dim r As Long
        Dim prodId As String
        Dim ws_prod As Worksheet
        Dim matchIdx As Variant
        
        On Error GoTo ErrHandler
        Application.EnableEvents = False
        
        Set ws_prod = Worksheets("Product Master")
        
        For Each Cell In TargetInterest
            r = Cell.Row
            prodId = Trim(Cell.Value)
            
            If prodId = "" Then
                ' Clear cells if Product ID is deleted
                Cells(r, 4).Value = "" ' Product Name
                Cells(r, 7).Value = "" ' Category
                Cells(r, 8).Value = "" ' Quantity
                Cells(r, 9).Value = "" ' Unit Price
            Else
                ' Search for Product ID in Product Master Column A (Columns(1))
                matchIdx = Application.Match(prodId, ws_prod.Columns(1), 0)
                
                If IsError(matchIdx) Then
                    MsgBox "Product ID '" & prodId & "' was not found in the Product Master registry!", vbExclamation + vbOKOnly, "Invalid Product ID"
                    
                    ' Clear the row details
                    Cell.Value = ""
                    Cells(r, 4).Value = ""
                    Cells(r, 7).Value = ""
                    Cells(r, 8).Value = ""
                    Cells(r, 9).Value = ""
                Else
                    ' Populate Product Name (Column B / 2), Category (Column C / 3) & Selling Price (Column I / 9)
                    Cells(r, 4).Value = ws_prod.Cells(matchIdx, 2).Value
                    Cells(r, 7).Value = ws_prod.Cells(matchIdx, 3).Value
                    Cells(r, 9).Value = ws_prod.Cells(matchIdx, 9).Value
                    
                    ' Default Quantity to 1 if currently empty or 0
                    If Val(Cells(r, 8).Value) <= 0 Then
                        Cells(r, 8).Value = 1
                    End If
                End If
            End If
        Next Cell
        
CleanExit:
        Application.EnableEvents = True
        Exit Sub
ErrHandler:
        MsgBox "An error occurred during auto-fill: " & Err.Description, vbCritical, "Auto-Fill Error"
        Resume CleanExit
    End If
End Sub
"""

    with open("VBA_PurchaseEntry.cls", "w", encoding="utf-8") as f: f.write(vba_purch)
    with open("VBA_SalesEntry.cls", "w", encoding="utf-8") as f: f.write(vba_sales)
    with open("VBA_Invoice.cls", "w", encoding="utf-8") as f: f.write(vba_inv)
    print("VBA script files generated in workspace.")

    # ----------------------------------------------------
    # 3. GENERATE GOOGLE APPS SCRIPT FILE
    # ----------------------------------------------------
    
    gas_code = """/**
 * Google Apps Script for Retail Inventory Management System Auto-fill
 * 
 * INSTRUCTIONS:
 * 1. Open your Google Sheet.
 * 2. Click on "Extensions" > "Apps Script".
 * 3. Delete any code in the editor, paste this entire script, and click "Save" (disk icon).
 * 4. This runs automatically on the "onEdit" event.
 */

function onEdit(e) {
  var range = e.range;
  var sheet = range.getSheet();
  var sheetName = sheet.getName();
  var col = range.getColumn();
  var row = range.getRow();
  
  var numRows = range.getNumRows();
  var numCols = range.getNumCols();
  
  // Ignore header edits and large sheet-clear operations
  if (row < 7 || numRows > 100) return;
  
  var ss = SpreadsheetApp.getActiveSpreadsheet();
  
  // 1. PURCHASE ENTRY AUTO-FILL
  if (sheetName === "Purchase Entry" && col === 3) {
    autoFillPurchase(sheet, range, row, numRows, ss);
  }
  
  // 2. SALES ENTRY AUTO-FILL
  else if (sheetName === "Sales Entry" && col === 3) {
    autoFillSales(sheet, range, row, numRows, ss);
  }
  
  // 3. INVOICE AUTO-FILL
  else if (sheetName === "Invoice" && col === 2 && row >= 11 && row <= 20) {
    autoFillInvoice(sheet, range, row, numRows, ss);
  }
}

function autoFillPurchase(sheet, range, startRow, numRows, ss) {
  var prodSheet = ss.getSheetByName("Product Master");
  var prodData = prodSheet.getRange("A7:W" + prodSheet.getLastRow()).getValues();
  
  var refSheet = ss.getSheetByName("ReferenceLists");
  var refData = refSheet.getRange("E2:F" + refSheet.getLastRow()).getValues();
  
  for (var i = 0; i < numRows; i++) {
    var r = startRow + i;
    var prodId = String(sheet.getRange(r, 3).getValue()).trim();
    
    if (prodId === "") {
      sheet.getRange(r, 4).setValue(""); // Product Name
      sheet.getRange(r, 5).setValue("").setDataValidation(null); // Supplier Name
      sheet.getRange(r, 7).setValue(""); // Cost Price
      continue;
    }
    
    // Find Product ID in Product Master
    var foundIndex = -1;
    for (var p = 0; p < prodData.length; p++) {
      if (String(prodData[p][0]).trim() === prodId) {
        foundIndex = p;
        break;
      }
    }
    
    if (foundIndex === -1) {
      SpreadsheetApp.getUi().alert("Product ID '" + prodId + "' not found in Product Master registry!");
      sheet.getRange(r, 3).setValue("");
      sheet.getRange(r, 4).setValue("");
      sheet.getRange(r, 5).setValue("").setDataValidation(null);
      sheet.getRange(r, 7).setValue("");
      continue;
    }
    
    // Extract info
    var prodName = prodData[foundIndex][1]; // B
    var costPrice = prodData[foundIndex][7]; // H
    var primarySup = prodData[foundIndex][4]; // E
    
    sheet.getRange(r, 4).setValue(prodName);
    sheet.getRange(r, 7).setValue(costPrice);
    
    // Query Supplier Mapping table
    var suppliers = [];
    for (var s = 0; s < refData.length; s++) {
      if (String(refData[s][0]).trim() === prodId) {
        suppliers.push(String(refData[s][1]).trim());
      }
    }
    
    // Fallback to primary supplier
    if (suppliers.length === 0) {
      suppliers.push(primarySup);
    }
    
    if (suppliers.length > 1) {
      // Create Dropdown Validation
      var rule = SpreadsheetApp.newDataValidation()
                               .requireValueInList(suppliers, true)
                               .setAllowInvalid(false)
                               .build();
      sheet.getRange(r, 5).setDataValidation(rule);
      sheet.getRange(r, 5).setValue(suppliers[0]);
    } else {
      // Single supplier: clear dropdown and set value
      sheet.getRange(r, 5).setDataValidation(null);
      sheet.getRange(r, 5).setValue(suppliers[0]);
    }
  }
}

function autoFillSales(sheet, range, startRow, numRows, ss) {
  var prodSheet = ss.getSheetByName("Product Master");
  var prodData = prodSheet.getRange("A7:W" + prodSheet.getLastRow()).getValues();
  
  for (var i = 0; i < numRows; i++) {
    var r = startRow + i;
    var prodId = String(sheet.getRange(r, 3).getValue()).trim();
    
    if (prodId === "") {
      sheet.getRange(r, 4).setValue(""); // Name
      sheet.getRange(r, 6).setValue(""); // Selling Price
      continue;
    }
    
    var foundIndex = -1;
    for (var p = 0; p < prodData.length; p++) {
      if (String(prodData[p][0]).trim() === prodId) {
        foundIndex = p;
        break;
      }
    }
    
    if (foundIndex === -1) {
      SpreadsheetApp.getUi().alert("Product ID '" + prodId + "' not found in Product Master registry!");
      sheet.getRange(r, 3).setValue("");
      sheet.getRange(r, 4).setValue("");
      sheet.getRange(r, 6).setValue("");
      continue;
    }
    
    sheet.getRange(r, 4).setValue(prodData[foundIndex][1]); // Name
    sheet.getRange(r, 6).setValue(prodData[foundIndex][8]); // Selling Price (I)
  }
}

function autoFillInvoice(sheet, range, startRow, numRows, ss) {
  var prodSheet = ss.getSheetByName("Product Master");
  var prodData = prodSheet.getRange("A7:W" + prodSheet.getLastRow()).getValues();
  
  for (var i = 0; i < numRows; i++) {
    var r = startRow + i;
    var prodId = String(sheet.getRange(r, 2).getValue()).trim(); // Invoice Product ID is Col 2 (B)
    
    if (prodId === "") {
      sheet.getRange(r, 4).setValue(""); // Name (D)
      sheet.getRange(r, 7).setValue(""); // Category (G)
      sheet.getRange(r, 8).setValue(""); // Qty (H)
      sheet.getRange(r, 9).setValue(""); // Unit Price (I)
      continue;
    }
    
    var foundIndex = -1;
    for (var p = 0; p < prodData.length; p++) {
      if (String(prodData[p][0]).trim() === prodId) {
        foundIndex = p;
        break;
      }
    }
    
    if (foundIndex === -1) {
      SpreadsheetApp.getUi().alert("Product ID '" + prodId + "' not found in Product Master registry!");
      sheet.getRange(r, 2).setValue("");
      sheet.getRange(r, 4).setValue("");
      sheet.getRange(r, 7).setValue("");
      sheet.getRange(r, 8).setValue("");
      sheet.getRange(r, 9).setValue("");
      continue;
    }
    
    sheet.getRange(r, 4).setValue(prodData[foundIndex][1]); // Name
    sheet.getRange(r, 7).setValue(prodData[foundIndex][2]); // Category
    sheet.getRange(r, 9).setValue(prodData[foundIndex][8]); // Price
    
    // Default Quantity to 1
    var currentQty = sheet.getRange(r, 8).getValue();
    if (currentQty === "" || currentQty <= 0) {
      sheet.getRange(r, 8).setValue(1);
    }
  }
}
"""

    with open("Google_Apps_Script.gs", "w", encoding="utf-8") as f: f.write(gas_code)
    print("Google Apps Script file generated in workspace.")

if __name__ == "__main__":
    add_autofill_assets()
